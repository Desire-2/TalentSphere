"""Server-side selectable-text PDF, DOCX, and TXT rendering."""

from __future__ import annotations

from dataclasses import dataclass
from html import escape
from io import BytesIO
import re
from typing import Any, Dict, Iterable, List, Tuple


class CareerDocumentRenderError(RuntimeError):
    pass


@dataclass
class RenderedCareerDocument:
    payload: bytes
    content_type: str
    filename: str
    page_count: int | None = None


def _safe_filename(value: str) -> str:
    cleaned = re.sub(r"[^A-Za-z0-9._-]+", "-", (value or "career-document").strip()).strip("-.")
    return cleaned[:100] or "career-document"


def _text(value: Any) -> str:
    return str(value or "").strip()


def _paragraphs(value: str) -> str:
    return "".join(f"<p>{escape(paragraph)}</p>" for paragraph in _text(value).split("\n") if paragraph.strip())


class CareerDocumentRenderer:
    """One content model rendered consistently for browser-independent exports."""

    def render(self, *, document_type: str, content: Dict[str, Any], settings: Dict[str, Any], title: str, format: str) -> RenderedCareerDocument:
        format = format.lower()
        base_name = _safe_filename(title)
        if format == "txt":
            return RenderedCareerDocument(self._to_txt(document_type, content).encode("utf-8"), "text/plain; charset=utf-8", f"{base_name}.txt")
        if format == "pdf":
            return self._to_pdf(document_type, content, settings, base_name)
        if format == "docx":
            return self._to_docx(document_type, content, settings, base_name)
        raise ValueError("Supported export formats are PDF, DOCX, and TXT.")

    def _to_pdf(self, document_type: str, content: Dict[str, Any], settings: Dict[str, Any], base_name: str) -> RenderedCareerDocument:
        try:
            from weasyprint import HTML
        except ImportError as exc:
            raise CareerDocumentRenderError("Server-side PDF rendering is unavailable.") from exc
        html = self._html(document_type, content, settings)
        try:
            payload = HTML(string=html).write_pdf()
        except Exception as exc:
            raise CareerDocumentRenderError("The document could not be rendered as PDF.") from exc
        page_count = self._pdf_page_count(payload)
        return RenderedCareerDocument(payload, "application/pdf", f"{base_name}.pdf", page_count)

    @staticmethod
    def _pdf_page_count(payload: bytes) -> int | None:
        try:
            from pypdf import PdfReader
            return len(PdfReader(BytesIO(payload)).pages)
        except Exception:
            return None

    def _to_docx(self, document_type: str, content: Dict[str, Any], settings: Dict[str, Any], base_name: str) -> RenderedCareerDocument:
        try:
            from docx import Document
            from docx.enum.text import WD_ALIGN_PARAGRAPH
            from docx.shared import Inches, Pt, RGBColor
        except ImportError as exc:
            raise CareerDocumentRenderError("DOCX export support is not installed on this server.") from exc
        document = Document()
        section = document.sections[0]
        section.top_margin = Inches(0.55)
        section.bottom_margin = Inches(0.55)
        section.left_margin = Inches(0.65)
        section.right_margin = Inches(0.65)
        normal = document.styles["Normal"]
        normal.font.name = "Aptos"
        normal.font.size = Pt(10)
        header = content.get("header", {})
        name = document.add_paragraph()
        name.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = name.add_run(_text(header.get("name")))
        run.bold = True
        run.font.size = Pt(20)
        title = document.add_paragraph()
        title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        title.add_run(_text(header.get("professional_title"))).italic = True
        contact = " · ".join(filter(None, [_text(header.get("email")), _text(header.get("phone")), _text(header.get("location"))] + [_text(link) for link in header.get("links", [])]))
        if contact:
            paragraph = document.add_paragraph(contact)
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.runs[0].font.size = Pt(8.5)
        if document_type == "COVER_LETTER":
            self._fill_cover_docx(document, content)
        else:
            self._fill_cv_docx(document, content)
        output = BytesIO()
        document.save(output)
        return RenderedCareerDocument(output.getvalue(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document", f"{base_name}.docx")

    @staticmethod
    def _fill_cv_docx(document, content: Dict[str, Any]) -> None:
        from docx.shared import Pt
        for section in content.get("sections", []):
            heading = document.add_paragraph()
            heading.paragraph_format.space_before = Pt(8)
            heading.paragraph_format.space_after = Pt(3)
            run = heading.add_run(_text(section.get("title")).upper())
            run.bold = True
            run.font.size = Pt(11)
            for entry in section.get("entries", []):
                title_line = " | ".join(filter(None, [_text(entry.get("heading")), _text(entry.get("subheading"))]))
                meta = " · ".join(filter(None, [_text(entry.get("dates")), _text(entry.get("location"))]))
                if title_line:
                    paragraph = document.add_paragraph()
                    paragraph.paragraph_format.space_after = Pt(0)
                    paragraph.add_run(title_line).bold = True
                    if meta:
                        paragraph.add_run(f"  {meta}").italic = True
                elif meta:
                    document.add_paragraph(meta)
                if _text(entry.get("text")):
                    document.add_paragraph(_text(entry.get("text")))
                for bullet in entry.get("bullets", []):
                    document.add_paragraph(_text(bullet.get("text")), style="List Bullet")

    @staticmethod
    def _fill_cover_docx(document, content: Dict[str, Any]) -> None:
        recipient = content.get("recipient", {})
        for value in (recipient.get("company"), recipient.get("role"), recipient.get("address")):
            if _text(value):
                document.add_paragraph(_text(value))
        if _text(content.get("subject")):
            paragraph = document.add_paragraph()
            paragraph.add_run(f"Re: {_text(content.get('subject'))}").bold = True
        if _text(recipient.get("salutation")):
            document.add_paragraph(_text(recipient.get("salutation")))
        for item in content.get("body", {}).get("paragraphs", []):
            document.add_paragraph(_text(item.get("text")))
        closing = _text(content.get("body", {}).get("closing")) or "Sincerely,"
        document.add_paragraph(closing)
        document.add_paragraph(_text(content.get("header", {}).get("name")))

    @staticmethod
    def _to_txt(document_type: str, content: Dict[str, Any]) -> str:
        header = content.get("header", {})
        header_lines = [
            _text(header.get("name")),
            _text(header.get("professional_title")),
            " | ".join(filter(None, [_text(header.get("email")), _text(header.get("phone")), _text(header.get("location"))])),
            " | ".join(_text(item) for item in header.get("links", []) if _text(item)),
        ]
        lines = [line for line in header_lines if line]
        if document_type == "COVER_LETTER":
            recipient = content.get("recipient", {})
            lines.extend(["", *[_text(recipient.get(key)) for key in ("company", "role", "address") if _text(recipient.get(key))]])
            if _text(content.get("subject")):
                lines.extend(["", f"Re: {_text(content.get('subject'))}"])
            if _text(recipient.get("salutation")):
                lines.extend(["", _text(recipient.get("salutation"))])
            lines.extend(["", *[_text(item.get("text")) for item in content.get("body", {}).get("paragraphs", []) if _text(item.get("text"))]])
            lines.extend(["", _text(content.get("body", {}).get("closing")) or "Sincerely,", _text(header.get("name"))])
            return "\n".join(lines).strip() + "\n"
        for section in content.get("sections", []):
            lines.extend(["", _text(section.get("title")).upper()])
            for entry in section.get("entries", []):
                title_line = " | ".join(filter(None, [_text(entry.get("heading")), _text(entry.get("subheading"))]))
                metadata = " · ".join(filter(None, [_text(entry.get("dates")), _text(entry.get("location"))]))
                if title_line:
                    lines.append(title_line)
                if metadata:
                    lines.append(metadata)
                if _text(entry.get("text")):
                    lines.append(_text(entry.get("text")))
                lines.extend([f"• {_text(bullet.get('text'))}" for bullet in entry.get("bullets", []) if _text(bullet.get("text"))])
        return "\n".join(lines).strip() + "\n"

    def _html(self, document_type: str, content: Dict[str, Any], settings: Dict[str, Any]) -> str:
        template_key = _text(settings.get("template_key") or settings.get("template") or "modern_professional")
        mode = _text(settings.get("document_mode") or "ats_safe")
        accent = _text(settings.get("accent_color") or "#0F766E")
        header = content.get("header", {})
        header_html = self._header_html(header, template_key, accent, mode)
        body_html = self._cover_html(content) if document_type == "COVER_LETTER" else self._cv_html(content)
        return f"""<!doctype html>
<html><head><meta charset=\"utf-8\"><style>
@page {{ size: A4; margin: 12mm 14mm 14mm; @bottom-center {{ content: counter(page); color: #777; font-size: 8pt; }} }}
* {{ box-sizing: border-box; }}
body {{ color: #18212f; font-family: Arial, Helvetica, sans-serif; font-size: 10pt; line-height: 1.42; margin: 0; }}
.document {{ max-width: 182mm; margin: 0 auto; }}
.header {{ border-bottom: 2px solid {accent}; padding: 0 0 8px; margin-bottom: 11px; }}
.header.executive {{ border-bottom-width: 1px; text-align: left; }}
.header.minimal {{ border-bottom-color: #333; }}
.header.creative_professional {{ background: #f7fbfa; padding: 10px 12px; border-left: 5px solid {accent}; border-bottom: 0; }}
.name {{ font-size: 22pt; letter-spacing: -.4pt; font-weight: 700; color: #102a43; margin: 0; }}
.role {{ color: {accent}; font-size: 10.5pt; font-weight: 600; margin-top: 2px; }}
.contact {{ color: #52616b; font-size: 8.5pt; margin-top: 5px; }}
.contact a {{ color: inherit; text-decoration: none; }}
.section {{ break-inside: avoid; margin: 0 0 10px; }}
.section-title {{ color: {accent}; font-size: 10.5pt; font-weight: 700; letter-spacing: .7pt; margin: 0 0 5px; padding-bottom: 2px; border-bottom: 1px solid #d8e1e7; text-transform: uppercase; }}
.entry {{ break-inside: avoid; margin: 0 0 7px; }}
.entry-head {{ display: flex; justify-content: space-between; gap: 12px; font-weight: 700; }}
.entry-sub {{ color: #3b4b5b; font-weight: 600; }}
.entry-meta {{ color: #62717e; font-size: 8.7pt; white-space: nowrap; }}
p {{ margin: 3px 0; }} ul {{ margin: 3px 0 0 15px; padding: 0; }} li {{ margin: 1.5px 0; }}
.cover {{ font-size: 10.5pt; line-height: 1.5; }} .recipient {{ margin-bottom: 13px; }} .subject {{ font-weight: 700; margin: 10px 0; }}
.ats-safe .header {{ border-bottom-color: #34495e; }} .ats-safe .section-title {{ color: #222; }}
</style></head><body><main class=\"document {escape(mode)}\">{header_html}{body_html}</main></body></html>"""

    @staticmethod
    def _header_html(header: Dict[str, Any], template_key: str, accent: str, mode: str) -> str:
        contact_items = [_text(header.get("email")), _text(header.get("phone")), _text(header.get("location"))] + [_text(link) for link in header.get("links", [])]
        contact = " · ".join(escape(item) for item in contact_items if item)
        return f"<header class=\"header {escape(template_key)}\"><h1 class=\"name\">{escape(_text(header.get('name')))}</h1><div class=\"role\">{escape(_text(header.get('professional_title')))}</div><div class=\"contact\">{contact}</div></header>"

    @staticmethod
    def _cv_html(content: Dict[str, Any]) -> str:
        pieces = []
        for section in content.get("sections", []):
            entries = []
            for entry in section.get("entries", []):
                heading = escape(_text(entry.get("heading")))
                subheading = escape(_text(entry.get("subheading")))
                dates = escape(_text(entry.get("dates")))
                location = escape(_text(entry.get("location")))
                head_left = "<span>" + heading + (f" <span class=\"entry-sub\">| {subheading}</span>" if subheading else "") + "</span>"
                head_right = " · ".join(item for item in (dates, location) if item)
                text = _paragraphs(_text(entry.get("text")))
                bullets = "".join(f"<li>{escape(_text(bullet.get('text')))}</li>" for bullet in entry.get("bullets", []) if _text(bullet.get("text")))
                entries.append(f"<article class=\"entry\"><div class=\"entry-head\">{head_left}<span class=\"entry-meta\">{head_right}</span></div>{text}{f'<ul>{bullets}</ul>' if bullets else ''}</article>")
            if entries:
                pieces.append(f"<section class=\"section\"><h2 class=\"section-title\">{escape(_text(section.get('title')))}</h2>{''.join(entries)}</section>")
        return "".join(pieces)

    @staticmethod
    def _cover_html(content: Dict[str, Any]) -> str:
        recipient = content.get("recipient", {})
        recipient_html = "<br>".join(escape(_text(recipient.get(key))) for key in ("company", "role", "address") if _text(recipient.get(key)))
        body = "".join(f"<p>{escape(_text(item.get('text')))}</p>" for item in content.get("body", {}).get("paragraphs", []) if _text(item.get("text")))
        subject = _text(content.get("subject"))
        salutation = _text(recipient.get("salutation"))
        closing = _text(content.get("body", {}).get("closing")) or "Sincerely,"
        signature = _text(content.get("header", {}).get("name"))
        subject_html = f'<div class="subject">Re: {escape(subject)}</div>' if subject else ''
        salutation_html = f'<p>{escape(salutation)}</p>' if salutation else ''
        return f"<article class=\"cover\"><div class=\"recipient\">{recipient_html}</div>{subject_html}{salutation_html}{body}<p>{escape(closing)}<br>{escape(signature)}</p></article>"
